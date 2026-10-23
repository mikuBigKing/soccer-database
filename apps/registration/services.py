"""报名业务服务层：流转、资料文件、Word/Excel 生成。"""
import io
import logging
import os
import uuid

from django.conf import settings
from django.utils import timezone

from apps.accounts.models import AthleteProfile, CoachProfile
from apps.common.codes import ResponseCode
from apps.common.exceptions import BusinessException
from apps.org.models import Team

from .models import Competition, Registration, RegistrationLog

business_logger = logging.getLogger('business')


def write_log(action: str, actor, registration: Registration = None, detail: dict = None) -> None:
    """报名日志统一入口。"""
    RegistrationLog.objects.create(
        registration=registration,
        action=action,
        actor=actor if getattr(actor, 'is_authenticated', False) else None,
        detail=detail or {},
    )
    business_logger.info(
        'registration log action=%s registration_id=%s actor=%s detail=%s',
        action, getattr(registration, 'pk', None), getattr(actor, 'pk', None), detail or {},
    )


class MaterialService:
    """报名资料文件：落私有目录，下载需鉴权（沿用 live.MatchFileService 口径）。"""

    SCAN_EXTS = {'.pdf', '.jpg', '.jpeg', '.png'}
    IMAGE_EXTS = {'.jpg', '.jpeg', '.png'}

    def save(self, relative_dir: str, uploaded_file, allowed_exts) -> tuple[str, str]:
        name = uploaded_file.name or 'file'
        ext = os.path.splitext(name)[1].lower()
        if ext not in allowed_exts:
            raise BusinessException(ResponseCode.FILE_TYPE_INVALID)
        if uploaded_file.size > settings.MAX_FILE_SIZE:
            raise BusinessException(ResponseCode.FILE_SIZE_EXCEEDED)

        absolute_dir = os.path.join(settings.PRIVATE_STORAGE_ROOT, relative_dir)
        os.makedirs(absolute_dir, exist_ok=True)
        stored_name = f'{uuid.uuid4().hex}{ext}'
        relative_path = os.path.join(relative_dir, stored_name)
        with open(os.path.join(settings.PRIVATE_STORAGE_ROOT, relative_path), 'wb') as target:
            for chunk in uploaded_file.chunks():
                target.write(chunk)
        return relative_path, name

    def absolute_path(self, relative_path: str) -> str:
        return os.path.join(settings.PRIVATE_STORAGE_ROOT, relative_path)

    def remove_quietly(self, relative_path: str) -> None:
        if not relative_path:
            return
        absolute = self.absolute_path(relative_path)
        if os.path.exists(absolute):
            os.remove(absolute)


def _get_registration(pk: int) -> Registration:
    obj = Registration.objects.filter(pk=pk).select_related('competition', 'team__school__city', 'team__coach').first()
    if obj is None:
        raise BusinessException(ResponseCode.NOT_FOUND, '报名单不存在', status_code=404)
    return obj


def _submitter_team_ids(user) -> set[int]:
    """提交方（教练/校级）可操作的球队 id 集合。"""
    if user.role == 'coach':
        return set(Team.objects.filter(coach=user).values_list('id', flat=True))
    if user.role == 'school':
        profile = user.role_profile()
        if profile and profile.school_id:
            return set(Team.objects.filter(school_id=profile.school_id).values_list('id', flat=True))
    return set()


def _city_id_of(user):
    profile = user.role_profile()
    return profile.city_id if profile and profile.city_id else None


class RegistrationService:
    """报名单流转服务。"""

    def __init__(self):
        self.materials = MaterialService()

    # ---------- 提交方 ----------

    def create_draft(self, user, competition_id: int, team_id: int) -> Registration:
        competition = Competition.objects.filter(pk=competition_id).first()
        if competition is None:
            raise BusinessException(ResponseCode.NOT_FOUND, '赛事不存在', status_code=404)
        if competition.status != Competition.Status.OPEN:
            raise BusinessException(ResponseCode.PARAM_INVALID, '该赛事报名已截止')
        if team_id not in _submitter_team_ids(user):
            raise BusinessException(ResponseCode.FORBIDDEN, '仅可为本校/本队球队报名')
        team = Team.objects.filter(pk=team_id).select_related('school').first()
        if team is None:
            raise BusinessException(ResponseCode.NOT_FOUND, '球队不存在', status_code=404)
        if Registration.objects.filter(competition=competition, team=team).exists():
            raise BusinessException(ResponseCode.DUPLICATE_SUBMIT, '该球队在此赛事下已有报名单')

        registration = Registration.objects.create(
            competition=competition, team=team, submitted_by=user,
        )
        write_log(RegistrationLog.Action.CREATE, user, registration, {'team_id': team_id})
        return registration

    def submit(self, user, registration: Registration) -> Registration:
        if registration.team_id not in _submitter_team_ids(user):
            raise BusinessException(ResponseCode.FORBIDDEN, '仅可提交本校/本队的报名单')
        if registration.status not in (Registration.Status.DRAFT, Registration.Status.REJECTED):
            raise BusinessException(ResponseCode.PARAM_INVALID, '当前状态不可提交')
        if not registration.stamped_form_path:
            raise BusinessException(ResponseCode.PARAM_INVALID, '请先上传盖章报名表扫描件')

        registration.status = Registration.Status.SUBMITTED
        registration.reject_level = None
        registration.reject_reason = None
        registration.submitted_at = timezone.now()
        registration.save(update_fields=['status', 'reject_level', 'reject_reason', 'submitted_at', 'updated_at'])
        write_log(RegistrationLog.Action.SUBMIT, user, registration)
        return registration

    def upload_stamped_form(self, user, registration: Registration, uploaded_file) -> Registration:
        if registration.team_id not in _submitter_team_ids(user):
            raise BusinessException(ResponseCode.FORBIDDEN, '仅可操作本校/本队的报名单')
        if registration.status not in (Registration.Status.DRAFT, Registration.Status.REJECTED):
            raise BusinessException(ResponseCode.PARAM_INVALID, '审核中或已通过的报名单不可更换扫描件')
        path, name = self.materials.save(f'registrations/{registration.pk}', uploaded_file, MaterialService.SCAN_EXTS)
        self.materials.remove_quietly(registration.stamped_form_path)
        registration.stamped_form_path = path
        registration.stamped_form_name = name
        registration.save(update_fields=['stamped_form_path', 'stamped_form_name', 'updated_at'])
        write_log(RegistrationLog.Action.MATERIAL_UPLOAD, user, registration, {'material': 'stamped_form', 'name': name})
        return registration

    # ---------- 市州端 ----------

    @staticmethod
    def _ensure_city_scope(user, registration: Registration) -> None:
        from apps.common.permissions import is_admin

        if is_admin(user):
            return
        if registration.city.pk != _city_id_of(user):
            raise BusinessException(ResponseCode.FORBIDDEN, '仅可操作本市州的报名单')

    def city_review(self, user, registration: Registration, approve: bool, reason: str = '') -> Registration:
        self._ensure_city_scope(user, registration)
        if registration.status != Registration.Status.SUBMITTED:
            raise BusinessException(ResponseCode.PARAM_INVALID, '仅待市州审核的报名单可审核')

        if approve:
            registration.status = Registration.Status.CITY_APPROVED
            registration.reject_level = None
            registration.reject_reason = None
            action = RegistrationLog.Action.CITY_APPROVE
        else:
            if not reason:
                raise BusinessException(ResponseCode.PARAM_INVALID, '驳回时必须填写驳回原因')
            registration.status = Registration.Status.REJECTED
            registration.reject_level = Registration.RejectLevel.CITY
            registration.reject_reason = reason
            action = RegistrationLog.Action.CITY_REJECT
        registration.city_reviewed_by = user.pk
        registration.city_reviewed_at = timezone.now()
        registration.save(update_fields=[
            'status', 'reject_level', 'reject_reason', 'city_reviewed_by', 'city_reviewed_at', 'updated_at',
        ])
        write_log(action, user, registration, {'reason': reason} if reason else {})
        return registration

    def report(self, user, registration: Registration) -> Registration:
        self._ensure_city_scope(user, registration)
        if registration.status != Registration.Status.CITY_APPROVED:
            raise BusinessException(ResponseCode.PARAM_INVALID, '仅市州通过的报名单可报送')
        registration.status = Registration.Status.REPORTED
        registration.reported_by = user.pk
        registration.reported_at = timezone.now()
        registration.save(update_fields=['status', 'reported_by', 'reported_at', 'updated_at'])
        write_log(RegistrationLog.Action.REPORT, user, registration)
        return registration

    def report_batch(self, user, ids: list[int]) -> dict:
        done, failed = [], []
        for pk in ids:
            try:
                self.report(user, _get_registration(pk))
                done.append(pk)
            except BusinessException as exc:
                failed.append({'id': pk, 'msg': exc.msg})
        return {'reported': done, 'failed': failed}

    # ---------- 省级端 ----------

    def province_review(self, user, registration: Registration, approve: bool, reason: str = '') -> Registration:
        if registration.status != Registration.Status.REPORTED:
            raise BusinessException(ResponseCode.PARAM_INVALID, '仅已报送的报名单可审核')

        if approve:
            registration.status = Registration.Status.APPROVED
            registration.reject_level = None
            registration.reject_reason = None
            action = RegistrationLog.Action.PROVINCE_APPROVE
        else:
            if not reason:
                raise BusinessException(ResponseCode.PARAM_INVALID, '驳回时必须填写驳回原因')
            registration.status = Registration.Status.REJECTED
            registration.reject_level = Registration.RejectLevel.PROVINCE
            registration.reject_reason = reason
            action = RegistrationLog.Action.PROVINCE_REJECT
        registration.province_reviewed_by = user.pk
        registration.province_reviewed_at = timezone.now()
        registration.save(update_fields=[
            'status', 'reject_level', 'reject_reason', 'province_reviewed_by', 'province_reviewed_at', 'updated_at',
        ])
        write_log(action, user, registration, {'reason': reason} if reason else {})
        return registration


def filter_registrations(queryset, params):
    """报名单列表通用过滤：keyword / status / competition_id / group_type / city_id。"""
    from django.db.models import Q

    keyword = params.get('keyword')
    if keyword:
        queryset = queryset.filter(
            Q(team__name__icontains=keyword)
            | Q(team__school__name__icontains=keyword)
            | Q(team__coach__name__icontains=keyword)
        )
    for param, lookup in (
        ('status', 'status'),
        ('competition_id', 'competition_id'),
        ('group_type', 'team__group_type'),
        ('city_id', 'team__school__city_id'),
    ):
        value = params.get(param)
        if value:
            queryset = queryset.filter(**{lookup: value})
    return queryset


def material_download(relative_path: str, filename: str, inline: bool = False):
    """私有目录文件鉴权下载/查看（调用方负责权限判断）。"""
    import mimetypes
    import os

    from django.http import FileResponse

    if not relative_path:
        raise BusinessException(ResponseCode.NOT_FOUND, '文件不存在', status_code=404)
    absolute = MaterialService().absolute_path(relative_path)
    if not os.path.exists(absolute):
        raise BusinessException(ResponseCode.NOT_FOUND, '文件不存在', status_code=404)
    mime = mimetypes.guess_type(filename or '')[0] or 'application/octet-stream'
    return FileResponse(open(absolute, 'rb'), as_attachment=not inline, filename=filename or 'file', content_type=mime)


class ExportService:
    """报名相关文件生成：Word 报名表 / Excel 报名情况导出。"""

    def registration_docx(self, registration: Registration) -> io.BytesIO:
        from docx import Document  # 延迟导入，未安装依赖时不影响迁移等命令

        team = registration.team
        school = team.school
        city = school.city
        coach = CoachProfile.objects.filter(user=team.coach).first()
        athletes = AthleteProfile.objects.filter(team=team).order_by('jersey_no', 'id')

        doc = Document()
        doc.add_heading(f'{registration.competition.name} 报名表', level=0)
        doc.add_paragraph(f'赛事年份：{registration.competition.year}')
        doc.add_paragraph(f'市州：{city.name}    学校：{school.name}')
        doc.add_paragraph(f'球队：{team.name}    组别：{team.get_group_type_display()}')
        if coach:
            doc.add_paragraph(f'教练员：{coach.name}    联系电话：{coach.phone or ""}    身份证号：{coach.id_card or ""}')
        doc.add_paragraph(f'报名状态：{registration.get_status_display()}    提交时间：'
                          f'{timezone.localtime(registration.submitted_at).strftime("%Y-%m-%d %H:%M") if registration.submitted_at else ""}')

        doc.add_heading('球员名单', level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        for cell, text in zip(table.rows[0].cells, ['球衣号', '姓名', '身份证号', '位置', '学籍材料']):
            cell.text = text
        for a in athletes:
            row = table.add_row().cells
            row[0].text = str(a.jersey_no or '')
            row[1].text = a.name
            row[2].text = a.id_card or ''
            row[3].text = a.position or ''
            row[4].text = '已上传' if a.enrollment_cert_path else '未上传'

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    def registrations_xlsx(self, queryset) -> io.BytesIO:
        from openpyxl import Workbook  # 延迟导入

        wb = Workbook()
        ws = wb.active
        ws.title = '报名情况'
        ws.append(['报名单号', '赛事', '年份', '市州', '学校', '球队', '组别', '状态',
                   '驳回方', '驳回原因', '提交时间', '市州审核时间', '报送时间', '省级审核时间'])

        def fmt(dt):
            return timezone.localtime(dt).strftime('%Y-%m-%d %H:%M') if dt else ''

        for r in queryset.select_related('competition', 'team__school__city'):
            ws.append([
                r.pk, r.competition.name, r.competition.year,
                r.team.school.city.name, r.team.school.name, r.team.name,
                r.team.get_group_type_display(), r.get_status_display(),
                r.get_reject_level_display() if r.reject_level else '', r.reject_reason or '',
                fmt(r.submitted_at), fmt(r.city_reviewed_at), fmt(r.reported_at), fmt(r.province_reviewed_at),
            ])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
